"""Read-only workbook audit; source workbook bytes are never modified."""

import json
import math
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
folder = ROOT / "data/external/figshare_30928694"
result = []
for path in sorted(folder.glob("*.xlsx")):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in book:
        rows = list(sheet.iter_rows(values_only=True))
        numbers = [
            v
            for row in rows
            for v in row
            if isinstance(v, (int, float)) and not isinstance(v, bool)
        ]
        labels = [
            {"row": i + 1, "column": j + 1, "text": v}
            for i, row in enumerate(rows)
            for j, v in enumerate(row)
            if isinstance(v, str) and not v.startswith("=")
        ]
        sheets.append(
            {
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "top_rows": [list(r[:12]) for r in rows[:7]],
                "text_labels": labels[:100],
                "numeric_cells": len(numbers),
                "nonfinite": sum(not math.isfinite(v) for v in numbers),
                "minimum": min(numbers) if numbers else None,
                "maximum": max(numbers) if numbers else None,
                "formula_cells": sum(
                    isinstance(v, str) and v.startswith("=")
                    for row in rows
                    for v in row
                ),
            }
        )
    result.append({"file": path.name, "sheets": sheets})
    book.close()
out = ROOT / "reports/external_cohort_audit.json"
out.write_text(
    json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
)
print(
    json.dumps(
        [
            {
                "file": f["file"],
                "sheets": len(f["sheets"]),
                "dimensions": sorted(
                    set((s["rows"], s["columns"]) for s in f["sheets"])
                ),
            }
            for f in result
        ]
    )
)
